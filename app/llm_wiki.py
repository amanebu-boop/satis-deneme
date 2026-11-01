"""Lightweight local LLM Wiki retrieval for Ersen Makina sales knowledge.

The index stores searchable excerpts locally in SQLite. Only the selected results
are sent to the OpenAI assistant when a user explicitly asks a relevant question.
"""
from __future__ import annotations

import hashlib
import json
import re
import time
from threading import Lock
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import APIRouter, File, HTTPException, Request, UploadFile

from app import db

router = APIRouter()
MAX_UPLOAD_BYTES = 8 * 1024 * 1024
MAX_EXTRACTED_CHARS = 200_000
_fts_available = False
_INDEX_REFRESH_SECONDS = 180
_last_refresh_at = 0.0
_refresh_lock = Lock()
_wiki_initialized = False

_ROOT = Path(__file__).resolve().parent.parent
_MAX_CONTENT = 12000
_TOKEN_RE = re.compile(r"[\w-]+", re.UNICODE)


def _tokens(value: str) -> list[str]:
    return [item.casefold() for item in _TOKEN_RE.findall(value or "") if len(item) > 1]


def _redact_meeting_contacts(value: str | None) -> str:
    text = value or ""
    text = re.sub(r"(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", "[e-posta gizlendi]", text)
    # Mask Turkish phone-number shapes while leaving ordinary short numeric values intact.
    phone_pattern = r"(?<!\w)(?:\+?90[ .-]?)?0?\(?[2-5]\d{2}\)?[ .-]?\d{3}[ .-]?\d{2}[ .-]?\d{2}(?!\w)"
    text = re.sub(phone_pattern, "[telefon gizlendi]", text)
    return re.sub(r"(?<!\d)\d{10,11}(?!\d)", "[vergi/kimlik numarası gizlendi]", text)


def init_wiki() -> None:
    global _wiki_initialized
    with db.connect() as connection:
        connection.execute(
            """CREATE TABLE IF NOT EXISTS llm_wiki_documents (
                id TEXT PRIMARY KEY,
                source_type TEXT NOT NULL,
                source_id TEXT NOT NULL,
                title TEXT NOT NULL,
                content TEXT NOT NULL,
                content_hash TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                UNIQUE(source_type, source_id)
            )"""
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS idx_llm_wiki_source ON llm_wiki_documents(source_type, source_id)"
        )
        global _fts_available
        try:
            connection.execute("CREATE VIRTUAL TABLE IF NOT EXISTS llm_wiki_fts USING fts5(title, content, tokenize='unicode61')")
            fts_count = connection.execute("SELECT COUNT(*) AS total FROM llm_wiki_fts").fetchone()["total"]
            _fts_available = True
            # Repair legacy orphan rows and backfill documents missing from the index.
            connection.execute(
                """DELETE FROM llm_wiki_fts
                   WHERE rowid NOT IN (
                       SELECT rowid FROM llm_wiki_documents
                       WHERE source_type != 'uploaded_document_meta'
                   )"""
            )
            connection.execute(
                """INSERT INTO llm_wiki_fts(rowid, title, content)
                   SELECT d.rowid, d.title, d.content
                   FROM llm_wiki_documents d
                   WHERE d.source_type != 'uploaded_document_meta'
                     AND NOT EXISTS (SELECT 1 FROM llm_wiki_fts f WHERE f.rowid=d.rowid)"""
            )
        except Exception:
            _fts_available = False
        _wiki_initialized = True


def _delete_wiki_documents(connection, where_sql: str, params: tuple[Any, ...]) -> None:
    """Delete matching documents and FTS rows in the same transaction."""
    rows = connection.execute(f"SELECT rowid FROM llm_wiki_documents WHERE {where_sql}", params).fetchall()
    if _fts_available and rows:
        connection.executemany("DELETE FROM llm_wiki_fts WHERE rowid=?", [(row["rowid"],) for row in rows])
    connection.execute(f"DELETE FROM llm_wiki_documents WHERE {where_sql}", params)


def _upsert(source_type: str, source_id: str, title: str, content: str) -> None:
    content = (content or "").strip()[:_MAX_CONTENT]
    if not content:
        return
    digest = hashlib.sha256((title + "\n" + content).encode("utf-8")).hexdigest()
    with db.connect() as connection:
        existing = connection.execute(
            "SELECT rowid, content_hash FROM llm_wiki_documents WHERE source_type=? AND source_id=?",
            (source_type, source_id),
        ).fetchone()
        changed = existing is None or existing["content_hash"] != digest
        connection.execute(
            """INSERT INTO llm_wiki_documents
               (id, source_type, source_id, title, content, content_hash, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
               ON CONFLICT(source_type, source_id) DO UPDATE SET
                 title=excluded.title, content=excluded.content,
                 content_hash=excluded.content_hash, updated_at=CURRENT_TIMESTAMP
               WHERE llm_wiki_documents.content_hash != excluded.content_hash""",
            (f"{source_type}:{source_id}", source_type, source_id, title, content, digest),
        )
        if _fts_available and source_type != "uploaded_document_meta" and changed:
            rowid = existing["rowid"] if existing else connection.execute(
                "SELECT rowid FROM llm_wiki_documents WHERE source_type=? AND source_id=?",
                (source_type, source_id),
            ).fetchone()["rowid"]
            connection.execute("DELETE FROM llm_wiki_fts WHERE rowid=?", (rowid,))
            connection.execute(
                "INSERT INTO llm_wiki_fts(rowid, title, content) VALUES (?, ?, ?)",
                (rowid, title, content),
            )


def refresh_wiki_index() -> None:
    """Refresh at most once per interval; serialize refreshes within this process."""
    global _last_refresh_at
    now = time.monotonic()
    if _last_refresh_at and now - _last_refresh_at < _INDEX_REFRESH_SECONDS:
        return
    if not _refresh_lock.acquire(blocking=False):
        return
    try:
        now = time.monotonic()
        if _last_refresh_at and now - _last_refresh_at < _INDEX_REFRESH_SECONDS:
            return
        _refresh_wiki_index_locked()
    finally:
        _refresh_lock.release()


def _refresh_wiki_index_locked() -> None:
    """Reconcile indexed sources. Schema/FTS initialization runs only at startup."""
    global _last_refresh_at
    if not _wiki_initialized:
        init_wiki()
    markdown_paths = [_ROOT / "AGENTS.md", _ROOT / "AGENT_ARCHITECTURE.md"]
    markdown_paths.extend(sorted((_ROOT / "agents").glob("*.md")))
    docs_root = _ROOT / "docs"
    if docs_root.exists():
        markdown_paths.extend(sorted(docs_root.rglob("*.md")))
    active_markdown_ids: set[str] = set()
    for path in markdown_paths:
        if not path.is_file():
            continue
        relative = path.relative_to(_ROOT).as_posix()
        active_markdown_ids.add(relative)
        try:
            _upsert("markdown", relative, relative, path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError):
            continue
    with db.connect() as connection:
        rows = connection.execute("SELECT source_id FROM llm_wiki_documents WHERE source_type='markdown'").fetchall()
        for row in rows:
            if row["source_id"] not in active_markdown_ids:
                _delete_wiki_documents(connection, "source_type='markdown' AND source_id=?", (row["source_id"],))

    products = db.list_products()
    product_by_id = {str(item["id"]): item for item in products}
    options_by_parent: dict[str, list[dict[str, Any]]] = {}
    for product in products:
        parent_id = product.get("parent_product_id")
        if parent_id:
            options_by_parent.setdefault(str(parent_id), []).append(product)
    active_ids: dict[str, set[str]] = {"product": set(), "company": set(), "meeting": set(), "quotation": set(), "follow_up": set()}
    for item in products:
        parent_id = item.get("parent_product_id")
        parent = product_by_id.get(str(parent_id)) if parent_id else None
        options = [
            {"sku": option.get("sku"), "name": option.get("name"),
             "description": option.get("description"), "currency": option.get("currency"),
             "unit_price": option.get("unit_price")}
            for option in options_by_parent.get(str(item["id"]), [])
        ]
        content = {
            "sku": item.get("sku"), "name": item.get("name"),
            "description": item.get("description"), "type": item.get("product_type"),
            "currency": item.get("currency"), "unit_price": item.get("unit_price"),
            "parent_product": parent.get("name") if parent else None,
            "options": options,
        }
        active_ids["product"].add(str(item["id"]))
        _upsert("product", str(item["id"]), str(item.get("name") or item.get("sku") or "Ürün"),
                json.dumps(content, ensure_ascii=False))

    for item in db.list_companies():
        sectors = item.get("sectors") or ([item.get("sector")] if item.get("sector") else [])
        content = {
            "name": item.get("name"), "sectors": sectors,
            "city": item.get("city"), "district": item.get("district"),
        }
        active_ids["company"].add(str(item["id"]))
        _upsert("company", str(item["id"]), str(item.get("name") or "Firma"),
                json.dumps(content, ensure_ascii=False))

    for item in db.list_meetings():
        content = {
            "company": item.get("company_name"), "date": item.get("meeting_date"),
            "type": item.get("meeting_type"), "topic": item.get("topic"),
            "notes": _redact_meeting_contacts(item.get("notes")),
            "products": [
                {"name": line.get("name"), "quantity": line.get("quantity")}
                for line in item.get("product_lines", [])
            ],
        }
        title = f"Görüşme · {item.get('company_name') or 'Firma'} · {item.get('meeting_date') or ''}"
        active_ids["meeting"].add(str(item["id"]))
        _upsert("meeting", str(item["id"]), title, json.dumps(content, ensure_ascii=False))

    with db.connect() as connection:
        quotation_rows = [dict(row) for row in connection.execute(
            """SELECT q.id, q.quote_number, q.currency, q.subtotal, q.tax_rate,
                      q.tax_amount, q.total, q.status, q.created_at, c.name AS company_name
               FROM quotations q LEFT JOIN companies c ON c.id = q.company_id
               ORDER BY q.created_at DESC, q.id DESC"""
        ).fetchall()]
        quote_lines = {
            str(quote["id"]): [dict(row) for row in connection.execute(
                """SELECT description, quantity, unit_price, line_total
                   FROM quotation_lines WHERE quotation_id=? ORDER BY rowid""",
                (quote["id"],),
            ).fetchall()]
            for quote in quotation_rows
        }
        follow_up_rows = [dict(row) for row in connection.execute(
            """SELECT t.id, t.title, t.due_date, t.status, c.name AS company_name
               FROM follow_up_tasks t JOIN companies c ON c.id=t.company_id
               ORDER BY t.due_date, t.id"""
        ).fetchall()]

    # Close the read connection before _upsert opens write connections.
    for quote in quotation_rows:
        content = {
            "quote_number": quote["quote_number"], "company": quote["company_name"],
            "currency": quote["currency"], "subtotal": quote["subtotal"],
            "tax_rate": quote["tax_rate"], "tax_amount": quote["tax_amount"],
            "total": quote["total"], "status": quote["status"],
            "created_at": quote["created_at"], "lines": quote_lines[str(quote["id"])],
        }
        active_ids["quotation"].add(str(quote["id"]))
        title = f"Teklif · {quote['quote_number'] or quote['id']} · {quote['company_name'] or 'Firma'}"
        _upsert("quotation", str(quote["id"]), title, json.dumps(content, ensure_ascii=False))

    for task in follow_up_rows:
        content = {
            "company": task["company_name"], "task": task["title"],
            "due_date": task["due_date"], "status": task["status"],
        }
        active_ids["follow_up"].add(str(task["id"]))
        _upsert("follow_up", str(task["id"]),
                f"Takip · {task['company_name']} · {task['title']}",
                json.dumps(content, ensure_ascii=False))

    # Deleted catalog/customer records must not remain searchable in the Wiki.
    with db.connect() as connection:
        for source_type, ids in active_ids.items():
            existing_ids = connection.execute(
                "SELECT source_id FROM llm_wiki_documents WHERE source_type=?",
                (source_type,),
            ).fetchall()
            for row in existing_ids:
                if row["source_id"] not in ids:
                    _delete_wiki_documents(connection, "source_type=? AND source_id=?", (source_type, row["source_id"]))


    _last_refresh_at = time.monotonic()


def _extract_document_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.casefold()
    if suffix in {".txt", ".md"}:
        return data.decode("utf-8-sig")
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        if len(reader.pages) > 100:
            raise HTTPException(status_code=413, detail="PDF bilgi belgesi en fazla 100 sayfa olabilir.")
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == ".docx":
        from docx import Document
        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    raise HTTPException(status_code=415, detail="Yalnızca PDF, DOCX, TXT veya MD dosyaları kabul edilir.")


@router.get("/api/wiki/documents")
def list_wiki_documents(request: Request) -> list[dict[str, Any]]:
    from app.main import _require_browser_session
    _require_browser_session(request)
    with db.connect() as connection:
        rows = connection.execute(
            """SELECT source_id, title, updated_at FROM llm_wiki_documents
               WHERE source_type='uploaded_document_meta' ORDER BY title"""
        ).fetchall()
    return [dict(row) for row in rows]


@router.post("/api/wiki/documents", status_code=201)
async def upload_wiki_document(request: Request, file: UploadFile = File(...)) -> dict[str, Any]:
    from app.main import _require_browser_session, _require_csrf
    _require_csrf(request, _require_browser_session(request))
    filename = Path(file.filename or "").name.strip()
    if not filename or Path(filename).suffix.casefold() not in {".pdf", ".docx", ".txt", ".md"}:
        raise HTTPException(status_code=415, detail="Yalnızca PDF, DOCX, TXT veya MD dosyaları kabul edilir.")
    data = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Bilgi belgesi en fazla 8 MB olabilir.")
    if not data:
        raise HTTPException(status_code=422, detail="Dosya boş.")
    try:
        content = _extract_document_text(filename, data).strip()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Belge metni okunamadı; dosya biçimini kontrol edin.") from exc
    if not content:
        raise HTTPException(status_code=422, detail="Belgede aranabilir metin bulunamadı. Taranmış PDF için OCR gerekir.")
    truncated = len(content) > MAX_EXTRACTED_CHARS
    if truncated:
        content = content[:MAX_EXTRACTED_CHARS]
    source_id = hashlib.sha256((filename + hashlib.sha256(data).hexdigest()).encode("utf-8")).hexdigest()[:32]
    # A newer upload with the same filename replaces the old indexed version.
    with db.connect() as connection:
        old_versions = connection.execute(
            """SELECT source_id FROM llm_wiki_documents
               WHERE source_type='uploaded_document_meta' AND lower(title)=lower(?)""",
            (filename,),
        ).fetchall()
        for old in old_versions:
            _delete_wiki_documents(connection, "source_type='uploaded_document' AND source_id LIKE ?", (str(old["source_id"]) + ":%",))
            _delete_wiki_documents(connection, "source_type='uploaded_document_meta' AND source_id=?", (old["source_id"],))
    _upsert("uploaded_document_meta", source_id, filename, json.dumps({"characters": len(content)}, ensure_ascii=False))
    # Chunk long manuals so later pages remain retrievable instead of indexing only
    # the first few thousand characters. Small overlap preserves split paragraphs.
    chunks = []
    position = 0
    while position < len(content):
        end = min(len(content), position + 1600)
        if end < len(content):
            boundary = content.rfind(" ", position + 1100, end)
            if boundary > position:
                end = boundary
        chunk = content[position:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(content):
            break
        position = max(position + 1, end - 180)
    with db.connect() as connection:
        _delete_wiki_documents(connection, "source_type='uploaded_document' AND source_id LIKE ?", (source_id + ":%",))
    for index, chunk in enumerate(chunks, start=1):
        _upsert("uploaded_document", f"{source_id}:{index:04d}",
                f"{filename} · Bölüm {index}", chunk)
    return {"source_id": source_id, "title": filename, "characters_indexed": len(content),
            "chunks_indexed": len(chunks), "truncated": truncated}


@router.delete("/api/wiki/documents/{source_id}")
def delete_wiki_document(source_id: str, request: Request) -> dict[str, bool]:
    from app.main import _require_browser_session, _require_csrf
    _require_csrf(request, _require_browser_session(request))
    with db.connect() as connection:
        exists = connection.execute(
            "SELECT 1 FROM llm_wiki_documents WHERE source_type='uploaded_document_meta' AND source_id=?",
            (source_id,),
        ).fetchone()
        if exists is None:
            raise HTTPException(status_code=404, detail="Bilgi belgesi bulunamadı.")
        _delete_wiki_documents(connection, "source_type='uploaded_document' AND source_id LIKE ?", (source_id + ":%",))
        _delete_wiki_documents(connection, "source_type='uploaded_document_meta' AND source_id=?", (source_id,))
    return {"deleted": True}


def search_wiki(query: str, limit: int = 5) -> dict[str, Any]:
    """Use local SQLite FTS5 when available, with lexical fallback; no external embeddings."""
    query = (query or "").strip()
    terms = list(dict.fromkeys(_tokens(query)))
    if not terms:
        return {"query": query, "results": [], "retrieval": "local_fts5" if _fts_available else "local_lexical"}
    refresh_wiki_index()
    safe_match = " OR ".join('"' + term.replace('"', '""') + '"' for term in terms)
    ranked = []
    retrieval = "local_lexical"
    with db.connect() as connection:
        if _fts_available:
            try:
                rows = connection.execute(
                    """SELECT d.source_type, d.source_id, d.title, d.content,
                              bm25(llm_wiki_fts, 5.0, 1.0) AS rank
                       FROM llm_wiki_fts JOIN llm_wiki_documents d ON d.rowid=llm_wiki_fts.rowid
                       WHERE llm_wiki_fts MATCH ? AND d.source_type != 'uploaded_document_meta'
                       ORDER BY rank ASC LIMIT ?""",
                    (safe_match, max(1, min(int(limit) * 4, 32))),
                ).fetchall()
                ranked = [(float(row["rank"]), row) for row in rows]
                retrieval = "local_fts5"
            except Exception:
                ranked = []
        if not ranked:
            clauses = " OR ".join("(title LIKE ? COLLATE NOCASE OR content LIKE ? COLLATE NOCASE)" for _ in terms)
            params = tuple(value for term in terms for value in (f"%{term}%", f"%{term}%"))
            rows = connection.execute(
                f"SELECT source_type, source_id, title, content FROM llm_wiki_documents WHERE source_type != 'uploaded_document_meta' AND ({clauses}) LIMIT 500",
                params,
            ).fetchall()
            for row in rows:
                title = str(row["title"] or "").casefold()
                content = str(row["content"] or "").casefold()
                score = sum((5 if term in title else 0) + content.count(term) for term in terms)
                if score:
                    ranked.append((-float(score), row))
            ranked.sort(key=lambda pair: (pair[0], pair[1]["source_type"], pair[1]["title"]))
    results = []
    for rank, row in ranked[:max(1, min(int(limit), 8))]:
        results.append({
            "source_type": row["source_type"], "title": row["title"],
            "content": str(row["content"])[:900],
            "relevance_score": round(-rank, 4) if retrieval == "local_fts5" else int(-rank),
        })
    return {"query": query, "retrieval": retrieval, "results": results}
