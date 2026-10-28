"""Private template storage and safe document generation for DOCX, XLSX and fillable PDF."""
from __future__ import annotations

import json
import os
import re
import zipfile
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from fastapi import HTTPException, UploadFile

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ZIP_MEMBERS = 2_000
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
ALLOWED_EXTENSIONS = {".docx", ".xlsx", ".pdf"}
PLACEHOLDER_RE = re.compile(r"\{\{\s*([a-zA-Z0-9_.-]{1,100})\s*\}\}")
SAFE_NAME_RE = re.compile(r"[^A-Za-z0-9._ -]+")


def storage_root() -> Path:
    root = Path(os.environ.get("SALES_DOCUMENTS_PATH", "data/documents")).resolve()
    root.mkdir(parents=True, exist_ok=True)
    return root


def _inside_root(path: Path) -> Path:
    root = storage_root()
    resolved = path.resolve()
    if resolved != root and root not in resolved.parents:
        raise HTTPException(status_code=400, detail="Invalid document path")
    return resolved


def safe_display_name(filename: str | None) -> str:
    raw = Path(filename or "template").name
    cleaned = SAFE_NAME_RE.sub("", raw).strip(" .")[:120]
    return cleaned or "template"


def _check_zip(data: bytes) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(BytesIO(data))
        members = archive.infolist()
        if len(members) > MAX_ZIP_MEMBERS:
            raise HTTPException(status_code=413, detail="Office file contains too many ZIP entries")
        if sum(item.file_size for item in members) > MAX_UNCOMPRESSED_BYTES:
            raise HTTPException(status_code=413, detail="Office file expands beyond the allowed size")
        if any(item.file_size > MAX_UNCOMPRESSED_BYTES for item in members):
            raise HTTPException(status_code=413, detail="Office file contains an oversized entry")
        if any("vbaproject.bin" in item.filename.lower() for item in members):
            raise HTTPException(status_code=415, detail="Macro-enabled Office files are not supported")
        return archive
    except zipfile.BadZipFile as exc:
        raise HTTPException(status_code=415, detail="Invalid Office document container") from exc


def validate_template(filename: str | None, data: bytes) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Only .docx, .xlsx and .pdf files are supported")
    if not data or len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413 if data else 400, detail="Template is empty or exceeds 10 MiB")
    if suffix in {".docx", ".xlsx"}:
        archive = _check_zip(data)
        names = set(archive.namelist())
        expected = {"word/document.xml"} if suffix == ".docx" else {"xl/workbook.xml"}
        if not expected.issubset(names):
            raise HTTPException(status_code=415, detail="File does not match the declared Office format")
        if any(name.lower().startswith("xl/externallinks/") for name in names):
            raise HTTPException(status_code=415, detail="Excel external links are not supported")
    elif not data.startswith(b"%PDF-"):
        raise HTTPException(status_code=415, detail="File does not match the PDF format")
    return suffix


def _docx_parts(document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
                for nested in cell.tables:
                    for nested_row in nested.rows:
                        for nested_cell in nested_row.cells:
                            for paragraph in nested_cell.paragraphs:
                                yield paragraph


def _replace_paragraph(paragraph, values: dict[str, str]) -> None:
    # Rebuild only when a placeholder is present. This supports placeholders split
    # across Word runs; formatting of the affected paragraph is normalized to its first run.
    original = "".join(run.text for run in paragraph.runs)
    if not PLACEHOLDER_RE.search(original):
        return
    replaced = PLACEHOLDER_RE.sub(lambda match: str(values.get(match.group(1), match.group(0))), original)
    if paragraph.runs:
        paragraph.runs[0].text = replaced
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replaced)


def discover_fields(suffix: str, data: bytes) -> list[str]:
    if suffix == ".docx":
        from docx import Document
        document = Document(BytesIO(data))
        text = [paragraph.text for paragraph in _docx_parts(document)]
        return sorted({match.group(1) for value in text for match in PLACEHOLDER_RE.finditer(value)})
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
        try:
            found = set()
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str):
                            found.update(match.group(1) for match in PLACEHOLDER_RE.finditer(cell.value))
            return sorted(found)
        finally:
            workbook.close()
    from pypdf import PdfReader
    reader = PdfReader(BytesIO(data), strict=True)
    if reader.is_encrypted:
        raise HTTPException(status_code=415, detail="Encrypted PDFs are not supported")
    fields = reader.get_fields() or {}
    return sorted(str(name) for name in fields)


def render_template(suffix: str, data: bytes, values: dict[str, str]) -> bytes:
    fields = set(discover_fields(suffix, data))
    missing = sorted(field for field in fields if field not in values)
    if missing:
        raise HTTPException(status_code=422, detail={"message": "Some template fields are unmapped", "missing_fields": missing})
    if suffix == ".docx":
        from docx import Document
        document = Document(BytesIO(data))
        for paragraph in _docx_parts(document):
            _replace_paragraph(paragraph, values)
        output = BytesIO()
        document.save(output)
        return output.getvalue()
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(data), data_only=False)
        try:
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and PLACEHOLDER_RE.search(cell.value):
                            cell.value = PLACEHOLDER_RE.sub(
                                lambda match: str(values.get(match.group(1), match.group(0))),
                                cell.value,
                            )
            output = BytesIO()
            workbook.save(output)
            return output.getvalue()
        finally:
            workbook.close()
    from pypdf import PdfReader, PdfWriter
    reader = PdfReader(BytesIO(data), strict=True)
    if reader.is_encrypted:
        raise HTTPException(status_code=415, detail="Encrypted PDFs are not supported")
    fields_in_pdf = reader.get_fields() or {}
    if not fields_in_pdf:
        raise HTTPException(status_code=422, detail="PDF has no fillable AcroForm fields; coordinate mapping is not supported yet")
    writer = PdfWriter()
    writer.clone_document_from_reader(reader)
    for page in writer.pages:
        writer.update_page_form_field_values(page, values, auto_regenerate=False)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def template_path(template_id: str) -> Path:
    if not re.fullmatch(r"[0-9a-f]{32}", template_id):
        raise HTTPException(status_code=404, detail="Template not found")
    path = _inside_root(storage_root() / f"{template_id}.source")
    if not path.is_file():
        raise HTTPException(status_code=404, detail="Template not found")
    return path


def metadata_path(template_id: str) -> Path:
    return _inside_root(storage_root() / f"{template_id}.json")


def read_metadata(template_id: str) -> dict:
    path = template_path(template_id)
    meta = metadata_path(template_id)
    if not meta.is_file():
        raise HTTPException(status_code=404, detail="Template not found")
    try:
        return json.loads(meta.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise HTTPException(status_code=500, detail="Template metadata is unavailable") from exc


async def save_upload(upload: UploadFile) -> dict:
    # Stream and cap bytes instead of trusting Content-Length or loading arbitrary uploads.
    chunks: list[bytes] = []
    total = 0
    while True:
        chunk = await upload.read(64 * 1024)
        if not chunk:
            break
        total += len(chunk)
        if total > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=413, detail="Template exceeds 10 MiB")
        chunks.append(chunk)
    data = b"".join(chunks)
    suffix = validate_template(upload.filename, data)
    fields = discover_fields(suffix, data)
    template_id = uuid4().hex
    name = safe_display_name(upload.filename)
    source = _inside_root(storage_root() / f"{template_id}.source")
    meta = metadata_path(template_id)
    source.write_bytes(data)
    try:
        meta.write_text(json.dumps({
            "id": template_id, "name": name, "format": suffix[1:],
            "fields": fields, "size": len(data),
        }, ensure_ascii=False), encoding="utf-8")
    except Exception:
        source.unlink(missing_ok=True)
        raise
    return {"id": template_id, "name": name, "format": suffix[1:], "fields": fields, "size": len(data)}


def generate_output(template_id: str, values: dict[str, str]) -> tuple[Path, str, str]:
    meta = read_metadata(template_id)
    source = template_path(template_id)
    data = source.read_bytes()
    suffix = "." + meta["format"]
    rendered = render_template(suffix, data, values)
    output_id = uuid4().hex
    output = _inside_root(storage_root() / f"{output_id}.output{suffix}")
    output.write_bytes(rendered)
    name = Path(meta["name"]).stem[:100] or "document"
    return output, name + suffix, suffix
